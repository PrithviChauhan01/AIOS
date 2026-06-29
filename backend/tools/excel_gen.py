"""Excel call-sheet export for leadgen prospect lists (openpyxl).

Turns the REAL Google Places rows into a working call sheet Sir can dial down:
styled header, frozen top row, and a status column with a fixed dropdown so the
sheet stays clean as he works it. Fields come verbatim from Places — nothing here
invents a value; email/notes are left blank for Sir to fill.
"""

import os
import re
import time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Status lifecycle for a cold-call sheet — the dropdown's allowed values.
STATUS_OPTIONS = [
    "new", "no_answer", "callback", "replied", "demo_booked",
    "won", "lost", "do_not_call",
]

# Column order + headers, exactly as Sir asked for the call sheet.
_COLUMNS = [
    ("name", "Name"),
    ("town", "Town"),
    ("phone", "Phone"),
    ("email", "Email"),
    ("website", "Website"),
    ("address", "Address"),
    ("rating", "Rating"),
    ("reviews", "Reviews"),
    ("status", "Status"),
    ("notes", "Notes"),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _slug(title: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", (title or "").lower()).strip("_")
    return s or "call_sheet"


def _town_from_address(address: str) -> str:
    """Best-effort town/city from a formatted address. Places formats like
    '123 Main St, Austin, TX 78701, USA' → 'Austin'. Falls back to blank rather
    than guessing — a blank cell is honest, a wrong town is not."""
    parts = [p.strip() for p in (address or "").split(",") if p.strip()]
    # Drop a trailing country and a 'STATE 12345' style region chunk, then take the
    # last remaining segment as the locality. Conservative: blank when unsure.
    if len(parts) >= 3:
        return parts[1] if len(parts) == 3 else parts[-3]
    return ""


def make_call_sheet(title: str, rows: list[dict], out_dir: str = "outputs") -> str:
    """Write the prospect rows to a styled .xlsx call sheet. Returns the path.

    rows are Places dicts {name, phone, website, address, rating, reviews}; email
    and notes are emitted blank, status defaults to 'new'."""
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{_slug(title)}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header row — styled.
    for col_idx, (_, header) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    # Data rows — values verbatim from Places; email/notes blank, status 'new'.
    for r in rows:
        record = {
            "name": r.get("name", ""),
            "town": _town_from_address(r.get("address", "")),
            "phone": r.get("phone", ""),
            "email": "",
            "website": r.get("website", ""),
            "address": r.get("address", ""),
            "rating": r.get("rating") if r.get("rating") is not None else "",
            "reviews": r.get("reviews", ""),
            "status": "new",
            "notes": "",
        }
        ws.append([record[key] for key, _ in _COLUMNS])

    # Borders on the data cells.
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(_COLUMNS)):
        for cell in row:
            cell.border = _BORDER

    # Status dropdown on every data row (col index of 'status').
    status_col = next(i for i, (k, _) in enumerate(_COLUMNS, start=1) if k == "status")
    status_letter = get_column_letter(status_col)
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(STATUS_OPTIONS) + '"',
        allow_blank=False,
        showDropDown=False,   # False here means: DO show the dropdown arrow (openpyxl quirk)
    )
    dv.error = "Pick a status from the list."
    dv.errorTitle = "Invalid status"
    ws.add_data_validation(dv)
    last_row = max(ws.max_row, 2)
    dv.add(f"{status_letter}2:{status_letter}{last_row}")

    # Frozen header + sensible column widths.
    ws.freeze_panes = "A2"
    _widths = {"Name": 28, "Town": 16, "Phone": 16, "Email": 24, "Website": 32,
               "Address": 40, "Rating": 8, "Reviews": 9, "Status": 14, "Notes": 30}
    for col_idx, (_, header) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = _widths.get(header, 16)

    wb.save(path)
    return path
